# =============================================================
# .xls, .xlsx, .csvのみを取り込みできる仕様
# 報告書A列の（取り込みキー）がDBの列名と一致している前提でrecordの変換辞書は本crawlerには無記載。ただ、DB登録側の関数が項目を受けられるように設定する必要がある
# =============================================================

import csv
import io
import sys
from datetime import date, datetime
from pathlib import Path

import xlrd
from openpyxl import load_workbook

from database import init_db, get_all_pages, insert_case

BASE_DIR = Path(__file__).resolve().parent
REPORT_DIR = BASE_DIR / "approved_reports"

SUPPORTED_EXTENSIONS = {".xlsx", ".xls", ".csv"}

HEADER_KEYS = {
    "（取り込み用キー）",
    "(取り込み用キー)",
    "（取り込みキー）",
    "(取り込みキー)",
}

DATE_KEYS = {
    "found_on",
    "occurred_on",
    "registered_on",
    "updated_at",
}

DATETIME_KEYS = {
    "manufactured_at",
}

SEARCH_TEXT_KEYS = {
    "defect_detail",
    "cause",
    "temp_action",
    "perm_action",
}

def read_rows(file_path: Path) -> list:
    extension = file_path.suffix.lower()

    if extension == ".xlsx":
        workbook = load_workbook(file_path, data_only=True)

        try:
            return list(
                workbook.active.iter_rows(
                    min_col=1,
                    max_col=3,
                    values_only=True,
                )
            )
        finally:
            workbook.close()
    
    if extension == ".xls":
        workbook = xlrd.open_workbook(str(file_path))

        try:
            sheet = workbook.sheet_by_index(0)
            rows = []

            for i in range(sheet.nrows):
                row = []

                for j in range(min(3, sheet.ncols)):
                    cell = sheet.cell(i, j)
                    value = cell.value

                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(
                            value,
                            workbook.datemode,
                        )

                    row.append(value)

                rows.append(row)

            return rows

        finally:
            workbook.release_resources()

    if extension == ".csv":
        raw = file_path.read_bytes()

        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp932")

        return list(csv.reader(io.StringIO(text, newline="")))

    raise ValueError(f"未対応の形式です: {extension}")

def fetch_data(file_path: Path) -> dict:
    """A列をキー、C列を値として辞書を返す。"""
    data = {}

    for row in read_rows(file_path):
        if not row or row[0] is None:
            continue

        key = str(row[0]).strip()

        if not key or key in HEADER_KEYS:
            continue

        # 同じキーを上書きしてしまうことを防ぐ
        if key in data:
            raise ValueError(f"取り込みキーが重複しています: {key}")

        value = row[2] if len(row) >= 3 else None

        if isinstance(value, str):
            value = value.strip()

        data[key] = None if value == "" else value

    return data

def normalize_data(data: dict) -> dict:
    """DBに渡す値を整える。"""
    record = dict(data)

    # 暫定対応：現状のDBの列名に合わせる
    # チームでstatusに統一したら、この変換は削除する
    if "status" in record:
        record["status_id"] = record.pop("status")

    for key in DATE_KEYS:
        value = record.get(key)

        if isinstance(value, (datetime, date)):
            record[key] = value.strftime("%Y-%m-%d")

    for key in DATETIME_KEYS:
        value = record.get(key)

        if isinstance(value, (datetime, date)):
            record[key] = value.strftime("%Y-%m-%d %H:%M")

    quantity = record.get("defect_qty")

    if quantity is not None:
        number = float(quantity)

        if not number.is_integer() or number < 0:
            raise ValueError("不良数量は0以上の整数にしてください")

        record["defect_qty"] = int(number)

    #👇現状のdatabase.pyのjoinがNoneで失敗しないための暫定対応です
    for key in SEARCH_TEXT_KEYS:
        value = record.get(key)
        record[key] = "" if value is None else str(value)

    part_no = record.get("part_no")

    if part_no is None or not str(part_no).strip():
        raise ValueError("品番（part_no）が未入力です")

    record["part_no"] = str(part_no).strip().upper()

    #database.pyの直接参照でkeyErrorにならないようにする
    #製品名が空欄なら、現在マスタは補完せずNULLで保存
    #👇
    record.setdefault("product_name", None)

    #数値セルのロット番号を文字列にする
    #Excelで失われた先頭のゼロは復元できない
    lot_no = record.get("lot_no")

    if lot_no is not None:
        if isinstance(lot_no, float) and lot_no.is_integer():
            lot_no = int(lot_no)

        record["lot_no"] = str(lot_no).strip()

    return record

def main() -> int:
    if not REPORT_DIR.is_dir():
        print(f"承認済みフォルダがありません: {REPORT_DIR}")
        return 1

    files = sorted(
        path
        for path in REPORT_DIR.iterdir()
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
        and not path.name.startswith("~$")
    )

    if not files:
        print(f"取り込み対象がありません: {REPORT_DIR}")
        return 0

    init_db()

    # 暫定対応：
    # report_fileに入っている報告書番号を重複判定に使う
    #👇
    existing_reports = {
        str(case["report_file"]).strip()
        for case in get_all_pages()
        if case.get("report_file")
    }

    inserted = 0
    skipped = 0
    failed = 0

    for file_path in files:
        try:
            data = fetch_data(file_path)
            record = normalize_data(data)

            report_id = str(record.get("report_file") or "").strip()

            if not report_id:
                raise ValueError(
                    "報告書番号（report_file）が未入力です"
                )

            record["report_file"] = report_id

            if report_id in existing_reports:
                skipped += 1
                print(f"スキップ: {file_path.name}（登録済み）")
                continue

            case_id = insert_case(record)
            existing_reports.add(report_id)
            inserted += 1

            print(
                f"登録成功: {file_path.name}"
                f" / case_id={case_id}"
            )

        except Exception as e:
            failed += 1
            print(
                f"登録失敗: {file_path.name} / {e}",
                file=sys.stderr,
            )

    print(
        f"結果: 登録 {inserted}件"
        f" / スキップ {skipped}件"
        f" / 失敗 {failed}件"
    )

    return 1 if failed else 0

if __name__ == "__main__":
    sys.exit(main())